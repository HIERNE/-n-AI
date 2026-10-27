import sys
from pathlib import Path
from datetime import datetime
from io import BytesIO
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

st.set_page_config(page_title="Hạch toán kế toán - InvoiceAI", page_icon="🧮", layout="wide")

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.storage.db_handler import get_confirmed_invoices, set_invoice_payment_status
from app.components.auth import render_account_menu, require_admin
from app.components.sidebar_nav import render_sidebar

admin_user = require_admin()
admin_user_id = int(admin_user["id"])
render_sidebar("accounting")
render_account_menu()

# --- Nạp style.css chuẩn hệ thống ---
css_path = Path(__file__).resolve().parents[1] / "style.css"
if css_path.exists():
    st.markdown(f"<style>{css_path.read_text(encoding='utf-8')}</style>", unsafe_allow_html=True)

# --- Header Banner Kỹ thuật ---
current_month_year = f"T{datetime.now().month}/{datetime.now().year}"

st.markdown(f"""
<div class="page-header">
    <div class="page-header-left">
        <div class="page-header-badge"><span>🧮</span> Quản lý tài chính & Báo cáo</div>
        <h1>Hạch toán kế toán</h1>
        <p>Tổng hợp số liệu hóa đơn đã xác nhận, theo dõi dòng tiền thanh toán và xuất dữ liệu báo cáo.</p>
    </div>
    <div style="display:flex; align-items:center; gap:12px; z-index:1;">
        <div class="header-period-pill">
            <span>Kỳ kế toán</span>
            {current_month_year}
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

payment_notice = st.session_state.pop("payment_status_notice", None)
if payment_notice:
    st.markdown(f"""
    <div class="success-notice">
        <span>✅</span> {payment_notice}
    </div>
    """, unsafe_allow_html=True)

invoices = get_confirmed_invoices(is_admin=True)

def format_amount(value):
    return f"{value:,.0f}".replace(",", ".") if value is not None else "—"

def create_excel_export(records):
    headers = [
        "STT", "ID", "Ngày lập", "Công ty bán", "MST bán", "Công ty mua",
        "MST mua", "Số hóa đơn", "Tiền hàng (đ)", "Tiền VAT (đ)",
        "Tổng TT (đ)", "Tình trạng thanh toán",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hóa đơn đã hạch toán"
    sheet.append(headers)

    for index, record in enumerate(records, start=1):
        sheet.append([
            index,
            record.get("id"),
            record.get("ngay_lap"),
            record.get("ten_cong_ty_ban"),
            record.get("ma_so_thue_ban"),
            record.get("ten_cong_ty_mua"),
            record.get("ma_so_thue_mua"),
            record.get("so_hoa_don"),
            record.get("tong_tien_hang") or 0,
            record.get("tien_thue_vat") or 0,
            record.get("tong_thanh_toan") or 0,
            "Đã thanh toán" if record.get("payment_status", "unpaid") == "paid" else "Chưa thanh toán",
        ])

    total_row = len(records) + 2
    sheet.cell(total_row, 1, "TỔNG CỘNG")
    for column in (9, 10, 11):
        letter = sheet.cell(1, column).column_letter
        sheet.cell(total_row, column, f"=SUM({letter}2:{letter}{total_row - 1})")

    header_fill = PatternFill("solid", fgColor="1F466D")
    total_fill = PatternFill("solid", fgColor="D9E5F5")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = total_fill
    for row in sheet.iter_rows(min_row=2, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = '#,##0'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{len(records) + 1}"
    widths = [8, 8, 14, 34, 18, 38, 18, 18, 18, 18, 18, 20]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

payment_status_values = {
    "Chưa thanh toán": "unpaid",
    "Đã thanh toán": "paid",
}
paid_invoices = [invoice for invoice in invoices if invoice.get("payment_status", "unpaid") == "paid"]
unpaid_invoices = [
    invoice for invoice in invoices if invoice.get("payment_status", "unpaid") == "unpaid"
]
paid_total = sum(invoice.get("tong_thanh_toan") or 0 for invoice in paid_invoices)
unpaid_total = sum(invoice.get("tong_thanh_toan") or 0 for invoice in unpaid_invoices)

# --- Thẻ số liệu tổng quan & Bộ lọc ---
col_filter, col_metrics = st.columns([1, 2.5])

with col_filter:
    st.markdown("""
    <div class="accounting-filter-card">
        <div style="font-size: 13px; font-weight: 700; color: #1E293B; margin-bottom: 6px;">🔍 Bộ lọc trạng thái</div>
    """, unsafe_allow_html=True)
    payment_filter = st.selectbox(
        "Lọc danh sách",
        ["Tất cả", "Chưa thanh toán", "Đã thanh toán"],
        label_visibility="collapsed",
    )
    st.markdown('<div style="font-size: 11.5px; color: #64748B; margin-top: 4px;">Áp dụng bộ lọc cho bảng dữ liệu và tệp Excel</div></div>', unsafe_allow_html=True)

with col_metrics:
    st.markdown(f"""
    <div class="metric-grid" style="margin-bottom:0;">
        <div class="metric-card blue" data-icon="📑">
            <div class="metric-label">Tổng đã xác nhận</div>
            <div class="metric-value">{len(invoices)}</div>
            <div class="metric-sub">Hóa đơn sẵn sàng hạch toán</div>
        </div>
        <div class="metric-card green" data-icon="💳">
            <div class="metric-label">Đã thanh toán</div>
            <div class="metric-value">{len(paid_invoices)}</div>
            <div class="metric-sub">{paid_total:,.0f} đ</div>
        </div>
        <div class="metric-card orange" data-icon="⏳">
            <div class="metric-label">Chưa thanh toán</div>
            <div class="metric-value">{len(unpaid_invoices)}</div>
            <div class="metric-sub">{unpaid_total:,.0f} đ</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

st.write("")

filtered_invoices = [
    invoice for invoice in invoices
    if payment_filter == "Tất cả"
    or invoice.get("payment_status", "unpaid") == payment_status_values[payment_filter]
]

if not invoices:
    st.markdown("""
    <div class="empty-state">
        <div class="empty-state-icon">📋</div>
        <h3>Chưa có dữ liệu hạch toán</h3>
        <p>Hiện chưa có hóa đơn nào được xác nhận để đưa vào danh sách hạch toán kế toán.</p>
    </div>
    """, unsafe_allow_html=True)
elif not filtered_invoices:
    st.markdown("""
    <div class="empty-state">
        <div class="empty-state-icon">🔍</div>
        <h3>Không tìm thấy kết quả</h3>
        <p>Không có hóa đơn nào phù hợp với tình trạng thanh toán đã chọn.</p>
    </div>
    """, unsafe_allow_html=True)
else:
    total_before_tax = sum(r.get("tong_tien_hang") or 0 for r in filtered_invoices)
    total_vat = sum(r.get("tien_thue_vat") or 0 for r in filtered_invoices)
    total_payment = sum(r.get("tong_thanh_toan") or 0 for r in filtered_invoices)

    st.markdown(f"""
    <div class="finance-summary-grid">
        <div class="finance-cell">
            <div class="finance-cell-label">Tiền hàng (Theo bộ lọc)</div>
            <div class="finance-cell-value">{total_before_tax:,.0f} đ</div>
        </div>
        <div class="finance-cell">
            <div class="finance-cell-label">Thuế VAT (Theo bộ lọc)</div>
            <div class="finance-cell-value">{total_vat:,.0f} đ</div>
        </div>
        <div class="finance-cell">
            <div class="finance-cell-label">Tổng thanh toán (Theo bộ lọc)</div>
            <div class="finance-cell-value highlight">{total_payment:,.0f} đ</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown(f"##### 📊 Chi tiết {len(filtered_invoices)} hóa đơn hạch toán")
    
    invoice_rows = [
        {
            "STT": index,
            "ID": record.get("id"),
            "Ngày lập": record.get("ngay_lap"),
            "Công ty bán": record.get("ten_cong_ty_ban"),
            "MST bán": record.get("ma_so_thue_ban"),
            "Công ty mua": record.get("ten_cong_ty_mua"),
            "MST mua": record.get("ma_so_thue_mua"),
            "Số hóa đơn": record.get("so_hoa_don") or "—",
            "Tiền hàng (đ)": format_amount(record.get("tong_tien_hang")),
            "Tiền VAT (đ)": format_amount(record.get("tien_thue_vat")),
            "Tổng TT (đ)": format_amount(record.get("tong_thanh_toan")),
            "Tình trạng thanh toán": (
                "Đã thanh toán"
                if record.get("payment_status", "unpaid") == "paid"
                else "Chưa thanh toán"
            ),
        }
        for index, record in enumerate(filtered_invoices, start=1)
    ]
    
    edited_invoices = st.data_editor(
        pd.DataFrame(invoice_rows),
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        disabled=[
            "STT", "ID", "Ngày lập", "Công ty bán", "MST bán",
            "Công ty mua", "MST mua", "Tiền hàng (đ)", "Tiền VAT (đ)", "Tổng TT (đ)",
        ],
        column_config={
            "Tình trạng thanh toán": st.column_config.SelectboxColumn(
                "Tình trạng thanh toán",
                options=["Chưa thanh toán", "Đã thanh toán"],
                required=True,
                width="medium",
            ),
        },
        key=f"invoice_payment_editor_{payment_filter}",
    )

    current_status_by_id = {
        invoice["id"]: invoice.get("payment_status", "unpaid")
        for invoice in filtered_invoices
    }
    updated_count = 0
    for edited_invoice in edited_invoices.to_dict(orient="records"):
        invoice_id = edited_invoice["ID"]
        selected_status = payment_status_values[edited_invoice["Tình trạng thanh toán"]]
        if selected_status != current_status_by_id[invoice_id]:
            set_invoice_payment_status(
                invoice_id,
                selected_status,
                is_admin=True,
                actor_user_id=admin_user_id,
            )
            updated_count += 1

    if updated_count:
        st.session_state["payment_status_notice"] = (
            f"Đã cập nhật tình trạng thanh toán cho {updated_count} hóa đơn."
        )
        st.rerun()

    # --- Khối Xuất file Excel ---
    col_dl_left, col_dl_right = st.columns([2, 1])
    with col_dl_left:
        st.markdown("""
        <div style="font-size: 13px; color: #475569; margin-top: 10px;">
            💡 <b>Mẹo kỹ thuật:</b> Tệp Excel xuất ra đã được định dạng tự động các công thức tính tổng (SUM), định dạng tiền tệ và bộ lọc dữ liệu sẵn cho phòng kế toán.
        </div>
        """, unsafe_allow_html=True)
    with col_dl_right:
        st.download_button(
            "⬇️ Xuất Excel đã hạch toán",
            data=create_excel_export(filtered_invoices),
            file_name="hoa_don_hach_toan_theo_bo_loc.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )

    st.divider()
    st.markdown("""
    <div class="info-bar">
        <span>⚙️</span>
        <span><b>Lưu ý:</b> Tính năng định khoản tự động (sinh bút toán Nợ/Có theo hệ thống tài khoản kế toán Việt Nam) đang được phát triển nâng cấp trong các phiên bản tiếp theo.</span>
    </div>
    """, unsafe_allow_html=True)